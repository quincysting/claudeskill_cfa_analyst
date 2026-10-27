"""xlsx_tool: every planted flaw in the Northwind DCF fixture is found; statements map to canonical items;
writing never touches the input."""
import hashlib

from pathlib import Path

import pytest

import xlsx_tool as xt

FIX = Path(__file__).resolve().parent / "fixtures"
DCF, FIN = FIX / "northwind_dcf.xlsx", FIX / "harbor_financials.xlsx"


@pytest.fixture(scope="module", autouse=True)
def fixtures():
    if not DCF.exists() or not FIN.exists():
        import runpy
        runpy.run_path(str(FIX / "make_fixtures.py"), run_name="__main__")


@pytest.fixture(scope="module")
def findings():
    return xt.audit(xt.Book(DCF))


@pytest.mark.parametrize("check, cell", [
    ("balance sheet does not balance", None),
    ("hardcoded value in formula row", "F5"),
    ("inconsistent formula", "E12"),
    ("inconsistent formula", "E4"),
    ("hardcoded number in formula", "E4"),
    ("FCFF from pre-tax EBIT", "C12"),
    ("terminal value discounted too few periods", "B18"),
    ("terminal growth above long-run nominal GDP", "B10"),
    ("WACC figures disagree", None),
    ("WACC weights at book value", None),
    ("terminal value dominates", "B18"),
])
def test_every_planted_flaw_is_found(findings, check, cell):
    hits = [f for f in findings if f["check"] == check and (cell is None or f["cell"] == cell)]
    assert hits, f"missing {check} {cell}: {[(f['check'], f['cell']) for f in findings]}"


def test_no_false_alarm_on_clean_rows(findings):
    flagged = {(f["sheet"], f["cell"]) for f in findings if f["check"] == "inconsistent formula"}
    assert flagged == {("DCF", "E12"), ("DCF", "E4")}


def test_findings_sorted_by_severity(findings):
    order = [xt.SEV[f["severity"]] for f in findings]
    assert order == sorted(order)


def test_extract_maps_statement_lines():
    tables = [t for t in xt.find_tables(xt.Book(FIN), "Financials")]
    cs = {}
    for t in tables:
        for k, v in xt.canonical_series(t).items():
            cs.setdefault(k, v)
    assert cs["revenue"] == [2000, 2200, 2420]
    assert cs["pretax_income"] == [200, 243, 295]
    assert cs["change_receivables"][1:] == [-80, -120]
    assert cs["short_term_debt"] == [50, 60, 120] and cs["long_term_debt"] == [500, 520, 560]
    assert cs["cfo"][1:] == pytest.approx([128.25, 98.25])


def test_canonical_labels_english_and_chinese():
    assert xt.canonical("Total revenues") == "revenue"
    assert xt.canonical("营业收入") == "revenue"
    assert xt.canonical("净利润") == "net_income"
    assert xt.canonical("Increase in inventory") == "change_inventory"
    assert xt.canonical("PV of terminal value") == "pv_terminal_value"
    assert xt.canonical("Footnote") is None


def test_review_and_patch_leave_the_input_untouched(tmp_path):
    before = hashlib.sha256(DCF.read_bytes()).hexdigest()
    fs = xt.audit(xt.Book(DCF))
    xt.write_review(DCF, fs, {"Enterprise value (corrected)": {"value": 1.0, "source": "test"}}, tmp_path / "r.xlsx")
    warnings = xt.patch(DCF, [{"sheet": "DCF", "cell": "F5", "value": "=F4*Assumptions!$B$4", "note": "restore"}],
                        tmp_path / "p.xlsx")
    assert hashlib.sha256(DCF.read_bytes()).hexdigest() == before
    assert warnings == []
    import openpyxl
    p = openpyxl.load_workbook(tmp_path / "p.xlsx")
    assert p["DCF"]["F5"].value == "=F4*Assumptions!$B$4" and "CFA Changes" in p.sheetnames
    r = openpyxl.load_workbook(tmp_path / "r.xlsx")
    assert r["Findings"].max_row == len(fs) + 1
    with pytest.raises(SystemExit):
        xt.patch(DCF, [], DCF)


def test_csv_input(tmp_path):
    p = tmp_path / "s.csv"
    p.write_text("Item,2022,2023\nRevenue,\"1,000\",\"1,100\"\nNet income,(50),80\n", encoding="utf-8")
    t = xt.find_tables(xt.Book(p), "s")[0]
    assert xt.canonical_series(t) == {"revenue": [1000.0, 1100.0], "net_income": [-50.0, 80.0]}


def test_ratios_from_workbook():
    out = xt.ratios(xt.Book(FIN))
    assert out["periods"] == ["2021", "2022", "2023"]
    r = out["by_period"]["2023"]
    assert r["interest_coverage"] == pytest.approx(340 / 45)
    assert r["current_ratio"] == pytest.approx(990 / 335)
    assert r["debt_to_equity"] == pytest.approx(680 / 1105)
    assert r["roe"] == pytest.approx(221.25 / ((980 + 1105) / 2))  # average equity: prior year is the beginning balance


def test_equity_rollforward_flags_unexplained_change():
    fs = [f for f in xt.audit(xt.Book(FIN)) if f["check"] == "equity roll-forward gap"]
    # fixture: 2022 equity 980 vs 840 + 182.25 - 64 = 958.25; 2023: 1105 vs 980 + 221.25 - 70 = 1131.25
    assert [f["cell"] for f in fs] == ["2022", "2023"]
    assert "gap 21.75" in fs[0]["detail"] and "gap -26.25" in fs[1]["detail"]


def test_wacc_label_match_is_exact():
    assert xt.canonical("WACC") == "wacc" and xt.canonical("WACC (computed)") == "wacc"
    assert xt.canonical("Discount rate") == "wacc"
    assert xt.canonical("WACC step") is None and xt.canonical("WACC sensitivity +0.5%") is None


def test_patch_log_keeps_formulas_as_text(tmp_path):
    import openpyxl
    xt.patch(DCF, [{"sheet": "DCF", "cell": "E4", "value": "=D4*(1+Assumptions!$B$3)", "note": "link growth"}],
             tmp_path / "p.xlsx")
    log = openpyxl.load_workbook(tmp_path / "p.xlsx")["CFA Changes"]
    row = [c for c in log[2]]
    assert row[2].data_type == "s" and row[3].data_type == "s" and row[3].value.startswith("=")
