"""Read, audit and annotate Excel workbooks for CFA-style analysis. Never modifies the input file.

  python xlsx_tool.py inspect FILE                         sheets, sizes, formulas, errors, links, hidden parts
  python xlsx_tool.py dump FILE [--sheet S] [--range A1:H40] [--max-rows 60]
                                                           grid with each cell's value and, for formulas, the formula
  python xlsx_tool.py extract FILE [--sheet S] [--json]    financial tables: periods x line items, labels mapped to
                                                           canonical names (revenue, net_income, total_assets, ...)
  python xlsx_tool.py audit FILE [--json]                  formula and model-integrity findings, most severe first
  python xlsx_tool.py ratios FILE [--sheet S]              per-period CFA ratio set (cfa_calc.financial_ratios) from
                                                           the extracted statements; prior year = beginning balances
  python xlsx_tool.py review FILE --findings F.json --out OUT.xlsx [--metrics M.json]
                                                           separate review workbook: findings + recomputed metrics
  python xlsx_tool.py patch FILE --changes C.json --out OUT.xlsx
                                                           corrected COPY: each change {sheet, cell, value, note}
                                                           is applied, commented and logged on a "CFA Changes" sheet

.xlsx/.xlsm are read twice (formulas, then cached values). .xls and .csv are read as values only.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
BENIGN_CONSTANTS = {0, 1, 2, 3, 4, 10, 12, 52, 100, 360, 365, 1000, 1e6, 0.5}
YEARISH = re.compile(r"^(FY|CY)?\s*'?(19|20)\d{2}\s*[AEFPB]?$|^\d{4}[AE]?$|^(Q[1-4]|H[12])\s*'?\d{2,4}|^(19|20)\d{2}[-/]"
                     r"|^Year\s*\d+$|^Y\d+$|^T\s*[+=]?\s*\d+$|^\d{1,2}$", re.I)

# canonical line items -> label patterns (lower-case, matched after normalisation); English and Chinese
CANON = {  # order matters: the first pattern that matches wins
    "change_receivables": r"(increase|decrease|change).*receivable|应收账款的(增加|减少)",
    "change_inventory": r"(increase|decrease|change).*inventor|存货的(增加|减少)",
    "change_payables": r"(increase|decrease|change).*payable|应付账款的(增加|减少)",
    "change_nwc": r"(increase|decrease|change).*(working capital|nwc)",
    "revenue": r"^(total )?(net )?(revenues?|sales|turnover)( revenue)?$|营业收入|营业总收入|销售收入|主营业务收入",
    "cogs": r"cost of (goods sold|sales|revenue)|^cogs$|营业成本|销售成本",
    "gross_profit": r"^gross (profit|margin)$|毛利",
    "sga": r"selling.*administrative|^sg&?a$|销售费用|管理费用",
    "rnd": r"research|^r&d$|研发",
    "ebitda": r"^(adj(usted)?\.? )?ebitda$",
    "depreciation": r"depreciation|amorti[sz]ation|^d&a$|折旧|摊销",
    "nopat": r"^nopat$|net operating profit after tax|ebit ?x? ?\(1 ?- ?t",
    "ebit": r"^ebit$|operating (income|profit)|income from operations|营业利润|息税前利润",
    "interest_expense": r"interest expense|^interest$|finance costs?|利息支出|财务费用",
    "pretax_income": r"(pre-?tax|before (income )?tax(es)?) (income|profit)|income before (income )?tax|^ebt$|利润总额|税前利润",
    "tax_rate": r"tax rate",
    "tax": r"(income )?tax(es)? (expense|provision)|^taxes?$|taxes on ebit|所得税",
    "net_income": r"^net (income|profit|earnings)|attributable to (shareholders|owners|common)|^净利润|归属于母公司",
    "eps": r"^(basic |diluted )?eps$|earnings per share|每股收益",
    "shares": r"shares outstanding|weighted average (number of )?shares|^shares$|股本|总股数",
    "dividends": r"dividends? (paid|declared)|^dividends?$|股利|分红",
    "cash": r"^cash( and (cash )?equivalents)?$|^cash & equivalents$|货币资金|现金及现金等价物$",
    "receivables": r"(accounts |trade )?receivables?|应收账款",
    "inventory": r"^inventor(y|ies)$|存货",
    "current_assets": r"total current assets|流动资产合计",
    "ppe": r"property,? plant|^pp&e|^(net )?fixed assets|固定资产",
    "total_assets": r"^total assets$|资产总计|总资产",
    "payables": r"(accounts |trade )?payables?$|应付账款",
    "current_liabilities": r"total current liabilities|流动负债合计",
    "short_term_debt": r"short-term (debt|borrowings)|current portion of (long-term )?debt|短期借款",
    "long_term_debt": r"long-term (debt|borrowings)|bonds payable|长期借款|应付债券",
    "total_debt": r"^total debt$|^(gross )?debt$|^borrowings$",
    "total_liabilities_and_equity": r"total liabilities and (shareholders'?|stockholders'? )?equity|负债和所有者权益总计",
    "total_liabilities": r"^total liabilities$|负债合计|总负债",
    "total_equity": r"total (shareholders'?|stockholders'?)? ?equity$|^total equity$|所有者权益合计|股东权益合计",
    "cfo": r"(net )?cash (flows? )?(from|provided by) operating|^cfo$|经营活动产生的现金流量净额",
    "capex": r"capital expenditures?|^capex$|purchases? of (property|pp&e|fixed)|购建固定资产",
    "cfi": r"(net )?cash (flows? )?(from|used in) investing|^cfi$|投资活动产生的现金流量净额",
    "cff": r"(net )?cash (flows? )?(from|used in) financing|^cff$|筹资活动产生的现金流量净额",
    "net_change_in_cash": r"net (increase|decrease|change).* in cash|现金及现金等价物净增加额",
    "fcf": r"free cash flow|^fcff$|^fcfe$|^ufcf$",
    "pv_terminal_value": r"(pv|present value) of (the )?terminal value|^pv of tv$",
    "terminal_value": r"terminal value|^tv$",
    "wacc": r"^(wacc|discount rate|weighted average cost of capital)( (computed|calculated|used|input|assumed|base|base case|nominal|real|estimate))?$",
    "terminal_growth": r"(terminal|perpetu(al|ity)|long[- ]term) growth|^g$",
    "enterprise_value": r"^enterprise value|^ev$",
    "cost_of_equity": r"cost of equity",
}
CANON_RE = {k: re.compile(v, re.I) for k, v in CANON.items()}


def norm_label(s) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[\(\)\[\]:,*]|\bin (millions|thousands)\b", " ", str(s))).strip().lower()


def canonical(label) -> str | None:
    n = norm_label(label)
    return next((k for k, rx in CANON_RE.items() if rx.search(n)), None) if n else None


# --- loading -----------------------------------------------------------------------------------
class Book:
    """Both views of a workbook: formulas (or None) and cached values, per sheet, as {(row, col): x}."""

    def __init__(self, path: Path):
        self.path, self.sheets, self.hidden, self.formulas, self.values = path, [], {}, {}, {}
        self.dims, self.names, self.parts, self.hidden_rows = {}, [], [], defaultdict(list)
        ext = path.suffix.lower()
        if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            self._load_openpyxl()
        elif ext == ".xls":
            self._load_xls()
        elif ext in (".csv", ".tsv"):
            self._load_csv("\t" if ext == ".tsv" else ",")
        else:
            raise SystemExit(f"unsupported file type {ext}")

    def _load_openpyxl(self):
        import openpyxl
        wb_f = openpyxl.load_workbook(self.path, data_only=False, read_only=False)
        wb_v = openpyxl.load_workbook(self.path, data_only=True, read_only=False)
        self.names = sorted(wb_f.defined_names.keys()) if hasattr(wb_f.defined_names, "keys") else []
        with zipfile.ZipFile(self.path) as z:
            names = z.namelist()
        self.parts = sorted({p.split("/")[1] for p in names if p.startswith("xl/") and p.count("/") >= 2
                             and p.split("/")[1] in ("charts", "pivotTables", "externalLinks", "drawings", "media")})
        for ws in wb_f.worksheets:
            self.sheets.append(ws.title)
            self.hidden[ws.title] = ws.sheet_state != "visible"
            self.dims[ws.title] = ws.dimensions
            f, v = {}, {}
            wv = wb_v[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    if isinstance(c.value, str) and c.value.startswith("="):
                        f[(c.row, c.column)] = c.value
                    elif type(c.value).__name__ == "ArrayFormula":
                        f[(c.row, c.column)] = "=" + str(getattr(c.value, "text", "")).lstrip("=")
                    v[(c.row, c.column)] = wv.cell(c.row, c.column).value if (c.row, c.column) in f else c.value
            self.formulas[ws.title], self.values[ws.title] = f, v
            self.hidden_rows[ws.title] = [r for r, d in ws.row_dimensions.items() if d.hidden]

    def _load_xls(self):
        import xlrd
        wb = xlrd.open_workbook(str(self.path))
        for sh in wb.sheets():
            self.sheets.append(sh.name)
            self.hidden[sh.name] = bool(sh.visibility)
            vals = {}
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    x = sh.cell_value(r, c)
                    if x not in ("", None):
                        vals[(r + 1, c + 1)] = x
            self.values[sh.name], self.formulas[sh.name] = vals, {}
            self.dims[sh.name] = f"A1:{col_letter(sh.ncols)}{sh.nrows}"

    def _load_csv(self, delim: str):
        rows = list(csv.reader(self.path.read_text(encoding="utf-8-sig").splitlines(), delimiter=delim))
        vals = {}
        for r, row in enumerate(rows, 1):
            for c, x in enumerate(row, 1):
                if x.strip():
                    vals[(r, c)] = to_number(x)
        name = self.path.stem
        self.sheets, self.values, self.formulas, self.hidden = [name], {name: vals}, {name: {}}, {name: False}
        self.dims[name] = f"A1:{col_letter(max((c for _, c in vals), default=1))}{len(rows)}"


def to_number(x):
    s = str(x).strip().replace(",", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").rstrip("%")
    try:
        v = float(s)
    except ValueError:
        return x
    v = -v if neg else v
    return v / 100 if str(x).strip().endswith("%") else v


def col_letter(c: int) -> str:
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def ref(r: int, c: int) -> str:
    return f"{col_letter(c)}{r}"


def is_num(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


# --- inspect / dump ------------------------------------------------------------------------------
def inspect_book(b: Book) -> dict:
    out = {"file": str(b.path), "sheets": [], "defined_names": b.names[:50], "special_parts": b.parts}
    for s in b.sheets:
        vals, f = b.values[s], b.formulas[s]
        errs = [ref(*k) for k, v in vals.items() if isinstance(v, str) and v.strip() in ERRORS]
        uncached = [ref(*k) for k in f if vals.get(k) is None]
        out["sheets"].append({"name": s, "hidden": b.hidden[s], "range": b.dims[s], "cells": len(vals) + len(
            [k for k in f if k not in vals]), "formulas": len(f), "error_cells": errs[:20], "n_errors": len(errs),
            "formulas_without_cached_value": len(uncached), "hidden_rows": b.hidden_rows.get(s, [])[:20],
            "external_refs": sum(1 for x in f.values() if "[" in x)})
    return out


def parse_range(rng: str | None) -> tuple[int, int, int, int] | None:
    if not rng:
        return None
    from openpyxl.utils import range_boundaries
    c1, r1, c2, r2 = range_boundaries(rng)
    if c1 is None or r1 is None or c2 is None or r2 is None:
        raise SystemExit(f"range {rng!r} must name both corners, e.g. A1:H40")
    return r1, c1, r2, c2


def dump(b: Book, sheet: str, rng: str | None, max_rows: int) -> str:
    vals, f = b.values[sheet], b.formulas[sheet]
    keys = set(vals) | set(f)
    if not keys:
        return f"(sheet {sheet} is empty)"
    bounds = parse_range(rng) or (min(r for r, _ in keys), min(c for _, c in keys),
                                 max(r for r, _ in keys), max(c for _, c in keys))
    r1, c1, r2, c2 = bounds
    r2 = min(r2, r1 + max_rows - 1)
    lines = [f"## {sheet}  {ref(r1, c1)}:{ref(r2, c2)}", ""]
    for r in range(r1, r2 + 1):
        cells = []
        for c in range(c1, c2 + 1):
            v, fx = vals.get((r, c)), f.get((r, c))
            if v is None and fx is None:
                continue
            shown = f"{v:.6g}" if isinstance(v, float) else str(v)
            cells.append(f"{ref(r, c)}={shown}" + (f" [{fx}]" if fx else ""))
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)


# --- extract financial tables ----------------------------------------------------------------------
def find_tables(b: Book, sheet: str) -> list[dict]:
    """A table = a header row with >= 2 period labels, a label column to its left, numeric rows below."""
    vals = b.values[sheet]
    rows = defaultdict(dict)
    for (r, c), v in vals.items():
        rows[r][c] = v
    tables, used = [], set()
    for r in sorted(rows):
        periods = {c: str(v).strip() for c, v in rows[r].items()
                   if (isinstance(v, str) and YEARISH.match(v.strip())) or (is_num(v) and 1990 <= float(v) <= 2100 and float(v).is_integer())
                   or type(v).__name__ in ("datetime", "date")}
        if len(periods) < 2 or r in used:
            continue
        pcols = sorted(periods)
        label_col = None
        items = {}
        for rr in range(r + 1, max(rows) + 1):
            row = rows.get(rr, {})
            labels = [(c, v) for c, v in row.items() if isinstance(v, str) and c < pcols[0] and v.strip()]
            nums = {c: row[c] for c in pcols if is_num(row.get(c))}
            if not row:
                continue
            if any(isinstance(v, str) and YEARISH.match(v.strip()) for c, v in row.items() if c in pcols) and \
                    sum(1 for c in pcols if isinstance(row.get(c), str)) >= 2:
                break  # next table's header
            if labels and len(nums) == 1 and len(pcols) >= 3 and pcols[0] in nums and \
                    not any((rr, c) in b.formulas[sheet] for c in pcols[1:]):
                continue  # a single value under the first period: a summary input (find_inputs), not a series
            if labels and nums:
                label_col = label_col or labels[-1][0]
                lab = labels[-1][1].strip()
                key = lab if lab not in items else f"{lab} (row {rr})"
                items[key] = {"row": rr, "canonical": canonical(lab),
                              "values": [row.get(c) if is_num(row.get(c)) else None for c in pcols],
                              "formula_cells": sum(1 for c in pcols if (rr, c) in b.formulas[sheet])}
                used.add(rr)
        if len(items) >= 2:
            tables.append({"sheet": sheet, "header_row": r, "periods": [periods[c] for c in pcols],
                           "period_columns": [col_letter(c) for c in pcols], "items": items})
    return tables


def canonical_series(table: dict) -> dict[str, list]:
    out = {}
    for lab, it in table["items"].items():
        if it["canonical"] and it["canonical"] not in out:
            out[it["canonical"]] = it["values"]
    return out


def find_inputs(b: Book, sheet: str) -> list[dict]:
    """Single-value rows: a label followed within four columns by one number (assumptions, DCF summary lines)."""
    rows = defaultdict(dict)
    for (r, c), v in b.values[sheet].items():
        rows[r][c] = v
    out = []
    for r, row in sorted(rows.items()):
        nums = [(c, v) for c, v in sorted(row.items()) if is_num(v)]
        labs = [(c, v) for c, v in sorted(row.items()) if isinstance(v, str) and v.strip() and v.strip() not in ERRORS]
        if len(nums) != 1 or not labs or labs[0][0] >= nums[0][0] or nums[0][0] - labs[0][0] > 4:
            continue
        c, v = nums[0]
        out.append({"sheet": sheet, "cell": ref(r, c), "label": labs[0][1].strip(), "canonical": canonical(labs[0][1]),
                    "value": v, "formula": b.formulas[sheet].get((r, c))})
    return out


def scalar_checks(b: Book) -> list[dict]:
    """Valuation sanity on single-value inputs across the whole workbook."""
    inputs = [i for s in b.sheets for i in find_inputs(b, s)]
    by = defaultdict(list)
    for i in inputs:
        if i["canonical"]:
            by[i["canonical"]].append(i)
    out = []

    def add(sev, where, check, detail, fix):
        out.append({"severity": sev, "sheet": where["sheet"], "cell": where["cell"], "check": check, "detail": detail,
                    "fix": fix})

    w, g = (by.get("wacc") or [None])[0], (by.get("terminal_growth") or [None])[0]
    if w and g:
        if w["value"] > 1 or g["value"] > 1:
            add("medium", w, "rates look like whole numbers", f"WACC {w['value']}, g {g['value']}",
                "store rates as decimals (8% = 0.08)")
        elif g["value"] >= w["value"]:
            add("high", g, "terminal growth >= discount rate", f"g {g['value']:.2%} vs WACC {w['value']:.2%}",
                "g must be below the discount rate; the Gordon formula is undefined otherwise")
        elif g["value"] > 0.05:
            add("medium", g, "terminal growth above long-run nominal GDP",
                f"g {g['value']:.2%} with WACC {w['value']:.2%} (spread {w['value'] - g['value']:.2%})",
                "stable growth should not exceed the economy's long-run nominal growth; test 2-4%")
    waccs = by.get("wacc", [])
    if len({round(x["value"], 4) for x in waccs if x["value"] <= 1}) > 1:
        detail = "; ".join(f"{x['sheet']}!{x['cell']} {x['label']} = {x['value']:.2%}" + ("" if x["formula"] else " (typed)")
                           for x in waccs)
        add("high", waccs[0], "WACC figures disagree", detail,
            "link the discount rate the valuation uses to one documented WACC build")
    book = [i for i in inputs if re.search(r"book value", i["label"], re.I)]
    if waccs and book:
        add("medium", book[0], "WACC weights at book value", ", ".join(f"{i['label']} {i['value']:,}" for i in book),
            "use market-value (target) capital-structure weights (CFA L1/L2 cost of capital)")
    tv, ev = (by.get("pv_terminal_value") or [None])[0], (by.get("enterprise_value") or [None])[0]
    if tv and ev and ev["value"] and tv["value"] / ev["value"] > 0.75:
        add("low", tv, "terminal value dominates", f"PV of TV is {tv['value'] / ev['value']:.0%} of EV",
            "test sensitivity to g and WACC; lengthen the explicit forecast if the firm is not yet in steady state")
    return out


def equity_rollforward(b: Book) -> list[dict]:
    """Ending equity vs prior equity + net income - dividends, period by period, across the workbook's statements."""
    periods, cs = statement_series(b)
    out = []
    if not {"total_equity", "net_income"} <= cs.keys():
        return out
    eq, ni, dv = cs["total_equity"], cs["net_income"], cs.get("dividends", [0] * len(periods))
    for i in range(1, len(periods)):
        if None in (eq[i], eq[i - 1], ni[i]):
            continue
        expected = eq[i - 1] + ni[i] - abs(dv[i] or 0)
        gap = eq[i] - expected
        if abs(gap) > 0.005 * max(1.0, abs(eq[i])):
            out.append({"severity": "medium", "sheet": "", "cell": f"{periods[i]}", "check": "equity roll-forward gap",
                        "detail": f"{periods[i]}: equity {eq[i]:,.2f} vs prior {eq[i - 1]:,.2f} + net income "
                                  f"{ni[i]:,.2f} - dividends {abs(dv[i] or 0):,.2f} = {expected:,.2f} (gap {gap:,.2f})",
                        "fix": "explain it (buybacks/issuance, OCI, other distributions) or treat the statements as "
                               "incomplete; an unexplained gap is a data-quality red flag"})
    return out


def dcf_lineage(b: Book) -> list[dict]:
    """DCF formula logic: FCFF built from pre-tax EBIT, terminal value discounted over too few periods."""
    out = []
    for s in b.sheets:
        f = b.formulas[s]
        row_label = {}
        for (r, c), v in b.values[s].items():
            if isinstance(v, str) and r not in row_label and canonical(v):
                row_label[r] = canonical(v)
        fcf_rows = [r for r, k in row_label.items() if k == "fcf"]
        n_proj = 0
        for r in fcf_rows:
            cells = sorted(c for (rr, c) in f if rr == r)
            n_proj = max(n_proj, len(cells))
            if not cells:
                continue
            refs = {int(m.group(1)) for m in re.finditer(r"(?<![!$A-Z])\$?[A-Z]{1,3}\$?(\d+)", f[(r, cells[0])])}
            used = {row_label.get(x) for x in refs}
            if "ebit" in used and not ({"nopat", "tax"} & used) and "nopat" in row_label.values():
                out.append({"severity": "high", "sheet": s, "cell": ref(r, cells[0]), "check": "FCFF from pre-tax EBIT",
                            "detail": f"{f[(r, cells[0])]} starts from EBIT although a NOPAT row exists",
                            "fix": "FCFF = EBIT(1 - t) + D&A - capex - change in NWC (CFA L2 free cash flow valuation)"})
        for r, k in row_label.items():
            if k != "pv_terminal_value":
                continue
            for (rr, c), fx in f.items():
                if rr == r:
                    m = re.search(r"\^\s*\(?\s*(\d+)\s*\)?", fx)
                    if m and n_proj and int(m.group(1)) < n_proj:
                        out.append({"severity": "high", "sheet": s, "cell": ref(rr, c),
                                    "check": "terminal value discounted too few periods",
                                    "detail": f"{fx} discounts {m.group(1)} periods; the forecast has {n_proj}",
                                    "fix": "discount the terminal value from the end of the final forecast year"})
    return out


# --- audit -------------------------------------------------------------------------------------------
SEV = {"high": 0, "medium": 1, "low": 2, "info": 3}


def audit(b: Book) -> list[dict]:
    from openpyxl.formula import Tokenizer
    from openpyxl.formula.translate import Translator
    fs: list[dict] = []

    def add(sev, sheet, cell, check, detail, fix=""):
        fs.append({"severity": sev, "sheet": sheet, "cell": cell, "check": check, "detail": detail, "fix": fix})

    if "externalLinks" in b.parts:
        add("medium", "", "", "external links", "workbook links to other files; values may be stale",
            "break or document the links; confirm the linked sources")
    for s in b.sheets:
        vals, f = b.values[s], b.formulas[s]
        if b.hidden[s]:
            add("low", s, "", "hidden sheet", "hidden sheets can hold overrides or stale calculations", "review it")
        if b.hidden_rows.get(s):
            add("low", s, f"rows {b.hidden_rows[s][:8]}", "hidden rows", "hidden rows with content", "review them")
        for (r, c), v in vals.items():
            if isinstance(v, str) and v.strip() in ERRORS:
                add("high", s, ref(r, c), "formula error", f"cell shows {v.strip()}", "trace precedents and fix")
        for (r, c), fx in f.items():
            if "[" in fx and "]" in fx:
                add("medium", s, ref(r, c), "external reference", fx[:80], "replace with an input cell")
            try:
                toks = Tokenizer(fx).items
            except Exception:  # noqa: BLE001 - malformed formula text
                continue
            consts = []
            for t in toks:
                if t.type == "OPERAND" and t.subtype == "NUMBER":
                    try:
                        x = float(t.value)
                    except ValueError:
                        continue
                    if x not in BENIGN_CONSTANTS:
                        consts.append(t.value)
            if consts:
                add("medium", s, ref(r, c), "hardcoded number in formula", f"{fx[:80]} contains {', '.join(consts[:4])}",
                    "move the assumption to a labelled input cell and reference it")
        # inconsistent formulas along a row (projection periods) and pasted values in formula rows
        by_row = defaultdict(list)
        for (r, c) in set(f) | {k for k, v in vals.items() if is_num(v)}:
            by_row[r].append(c)
        for r, cols in by_row.items():
            cols.sort()
            fcols = [c for c in cols if (r, c) in f]
            if len(fcols) < 3:
                continue
            for i, c in enumerate(fcols):
                expected = []
                for n in (fcols[i - 1] if i else None, fcols[i + 1] if i + 1 < len(fcols) else None):
                    if n is not None and abs(n - c) == 1:
                        try:
                            expected.append(Translator(f[(r, n)], origin=ref(r, n)).translate_formula(ref(r, c)))
                        except Exception:  # noqa: BLE001
                            pass
                if len(expected) == 2 and expected[0] == expected[1] and expected[0] != f[(r, c)]:
                    add("high", s, ref(r, c), "inconsistent formula",
                        f"{f[(r, c)][:60]} breaks the row pattern {expected[0][:60]}", "make it consistent or document why")
            lo, hi = fcols[0], fcols[-1]
            for c in cols:
                if lo < c < hi and (r, c) not in f and is_num(vals.get((r, c))):
                    add("high", s, ref(r, c), "hardcoded value in formula row",
                        f"{vals[(r, c)]} typed over a row of formulas", "restore the formula or flag it as an override")
        # financial integrity on extracted tables
        for t in find_tables(b, s):
            fs.extend(integrity(t))
    fs.extend(scalar_checks(b))
    fs.extend(dcf_lineage(b))
    fs.extend(equity_rollforward(b))
    fs.sort(key=lambda x: (SEV[x["severity"]], x["sheet"], x["check"]))
    return fs


def _close(a, b, tol=0.005):
    return a is not None and b is not None and abs(a - b) <= tol * max(1.0, abs(a), abs(b))


def integrity(t: dict) -> list[dict]:
    """Accounting identities and valuation sanity on one extracted table."""
    cs, out, s = canonical_series(t), [], t["sheet"]
    per = t["periods"]

    def add(sev, check, detail, fix):
        out.append({"severity": sev, "sheet": s, "cell": f"table at row {t['header_row']}", "check": check,
                    "detail": detail, "fix": fix})

    n = len(per)
    if {"total_assets", "total_liabilities", "total_equity"} <= cs.keys():
        for i in range(n):
            a, l, e = cs["total_assets"][i], cs["total_liabilities"][i], cs["total_equity"][i]
            if None not in (a, l, e) and not _close(a, l + e):
                add("high", "balance sheet does not balance", f"{per[i]}: assets {a:,.2f} vs L+E {l + e:,.2f} "
                    f"(gap {a - l - e:,.2f})", "trace the gap before any other analysis")
    if {"cfo", "cfi", "cff", "net_change_in_cash"} <= cs.keys():
        for i in range(n):
            o, iv, fi, d = (cs[k][i] for k in ("cfo", "cfi", "cff", "net_change_in_cash"))
            if None not in (o, iv, fi, d) and not _close(o + iv + fi, d, 0.01):
                add("high", "cash flow does not sum", f"{per[i]}: CFO+CFI+CFF {o + iv + fi:,.2f} vs change "
                    f"{d:,.2f}", "check FX effects and sign conventions")
    if {"revenue", "cogs", "gross_profit"} <= cs.keys():
        for i in range(n):
            r, c, g = cs["revenue"][i], cs["cogs"][i], cs["gross_profit"][i]
            if None not in (r, c, g) and not (_close(r - abs(c), g) or _close(r + c, g)):
                add("medium", "gross profit does not tie", f"{per[i]}: revenue {r:,.2f}, COGS {c:,.2f}, GP {g:,.2f}",
                    "check the COGS sign and definition")
    if "revenue" in cs:
        rev = [x for x in cs["revenue"] if x is not None]
        g = [(b2 / a - 1) for a, b2 in zip(rev, rev[1:]) if a]
        if any(x > 1.0 for x in g):
            add("low", "revenue growth above 100%", ", ".join(f"{x:.0%}" for x in g), "justify or temper the forecast")
    w = next((x for x in cs.get("wacc", []) if x is not None), None)
    tg = next((x for x in cs.get("terminal_growth", []) if x is not None), None)
    if w is not None and tg is not None:
        if w > 1 or tg > 1:
            add("medium", "rates look like whole numbers", f"WACC {w}, g {tg}", "store rates as decimals (8% = 0.08)")
        elif tg >= w:
            add("high", "terminal growth >= discount rate", f"g {tg:.2%} vs WACC {w:.2%}: Gordon value undefined",
                "g must be below the discount rate (CFA L2 FCF valuation)")
        elif tg > 0.05:
            add("medium", "terminal growth above long-run nominal GDP", f"g {tg:.2%}",
                "stable growth should not exceed the economy's nominal growth rate")
    tv = next((x for x in cs.get("terminal_value", []) if x is not None), None)
    ev = next((x for x in cs.get("enterprise_value", []) if x is not None), None)
    if tv and ev and tv / ev > 0.75:
        add("low", "terminal value dominates", f"TV/EV = {tv / ev:.0%} (if TV is already discounted)",
            "test sensitivity to g and WACC; check the explicit-period length")
    tr = [x for x in cs.get("tax_rate", []) if x is not None]
    if any(x > 1 for x in tr):
        add("medium", "tax rate stored as whole number", str(tr[:5]), "store as a decimal")
    elif any(x < 0 or x > 0.5 for x in tr):
        add("low", "unusual tax rate", str(tr[:5]), "reconcile effective vs statutory rate")
    return out


# --- ratios: extracted statements -> cfa_calc.financial_ratios ------------------------------------------
def statement_series(b: Book, sheet: str | None = None) -> tuple[list[str], dict[str, list]]:
    """Canonical line items across the workbook's period tables (first occurrence wins), on the first table's periods."""
    periods, series = [], {}
    for s in ([sheet] if sheet else b.sheets):
        for t in find_tables(b, s):
            if not periods:
                periods = t["periods"]
            if len(t["periods"]) != len(periods):
                continue
            for k, v in canonical_series(t).items():
                series.setdefault(k, v)
    if "total_debt" not in series and ("short_term_debt" in series or "long_term_debt" in series):
        st, lt = series.get("short_term_debt", [0] * len(periods)), series.get("long_term_debt", [0] * len(periods))
        series["total_debt"] = [None if a is None and c is None else (a or 0) + (c or 0) for a, c in zip(st, lt)]
    return periods, series


def ratios(b: Book, sheet: str | None = None) -> dict:
    """Per-period CFA ratios from the workbook's statements; prior-period balances become beginning balances."""
    import inspect
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from cfa_calc import REGISTRY
    fn = REGISTRY["financial_ratios"].func
    params = set(inspect.signature(fn).parameters)
    periods, series = statement_series(b, sheet)
    out = {"periods": periods, "inputs_found": sorted(k for k in series if k in params), "by_period": {}}
    for i, p in enumerate(periods):
        kw = {k: series[k][i] for k in params if k in series and series[k][i] is not None}
        if i:
            for k in params:
                if k.endswith("_begin") and k[:-6] in series and series[k[:-6]][i - 1] is not None:
                    kw[k] = series[k[:-6]][i - 1]
        out["by_period"][p] = fn(**kw) if kw else {}
    return out


# --- writing: review workbook and patched copy ----------------------------------------------------------
def write_review(src: Path, findings: list[dict], metrics: dict | None, out: Path) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    if out.resolve() == src.resolve():
        raise SystemExit("refusing to overwrite the input workbook")
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0]
    ws.title = "Findings"
    head = ["#", "Severity", "Sheet", "Cell", "Check", "Detail", "Recommended fix", "CFA reference"]
    ws.append(head)
    for i, f in enumerate(findings, 1):
        ws.append([i, f.get("severity"), f.get("sheet"), f.get("cell"), f.get("check"), f.get("detail"),
                   f.get("fix"), f.get("cfa_reference", "")])
    fill = {"high": "F8CBAD", "medium": "FFE699", "low": "DDEBF7", "info": "EDEDED"}
    for row in ws.iter_rows(min_row=2):
        row[1].fill = PatternFill("solid", fgColor=fill.get(str(row[1].value), "FFFFFF"))
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (5, 10, 16, 14, 30, 60, 50, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    if metrics:
        m = wb.create_sheet("Recomputed")
        m.append(["Metric", "Value", "Source / formula"])
        for k, v in metrics.items():
            if isinstance(v, dict):
                m.append([k, v.get("value"), v.get("source", "")])
            else:
                m.append([k, v, ""])
        for c in m[1]:
            c.font = Font(bold=True)
        m.column_dimensions["A"].width, m.column_dimensions["C"].width = 40, 60
    meta = wb.create_sheet("About")
    meta.append(["Reviewed file", str(src)])
    meta.append(["Method", "cfa-analyst skill: xlsx_tool.py audit + cfa_calc recomputation"])
    wb.save(out)


def patch(src: Path, changes: list[dict], out: Path) -> list[str]:
    import openpyxl
    from openpyxl.comments import Comment
    if out.resolve() == src.resolve():
        raise SystemExit("refusing to overwrite the input workbook")
    with zipfile.ZipFile(src) as z:
        lost = sorted({p.split("/")[1] for p in z.namelist() if p.startswith("xl/") and p.count("/") >= 2
                       and p.split("/")[1] in ("charts", "pivotTables", "drawings", "media")})
    shutil.copyfile(src, out)
    wb = openpyxl.load_workbook(out, keep_vba=src.suffix.lower() == ".xlsm")
    log = wb["CFA Changes"] if "CFA Changes" in wb.sheetnames else wb.create_sheet("CFA Changes")
    if log.max_row == 1 and log.cell(1, 1).value is None:
        for j, h in enumerate(["Sheet", "Cell", "Old", "New", "Why"], 1):
            log.cell(1, j, h)
    for ch in changes:
        ws = wb[ch["sheet"]]
        cell = ws[ch["cell"]]
        old = cell.value
        cell.value = ch["value"]
        cell.comment = Comment(f"CFA review: {ch.get('note', '')} (was: {old})", "cfa-analyst")
        log.append([ch["sheet"], ch["cell"], str(old), str(ch["value"]), ch.get("note", "")])
        for c in log[log.max_row]:  # old/new formulas are a record, not live formulas on the log sheet
            c.data_type = "s"
    wb.calculation.fullCalcOnLoad = True  # Excel recomputes every formula on open
    wb.save(out)
    return [f"openpyxl does not keep {', '.join(lost)} from the original; they are missing in the copy"] if lost else []


def main(argv: list[str]) -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("inspect", "dump", "extract", "audit", "ratios", "review", "patch"):
        p = sub.add_parser(name)
        p.add_argument("file")
        if name in ("dump", "extract", "ratios"):
            p.add_argument("--sheet")
        if name == "dump":
            p.add_argument("--range")
            p.add_argument("--max-rows", type=int, default=60)
        if name in ("extract", "audit"):
            p.add_argument("--json", action="store_true")
        if name == "review":
            p.add_argument("--findings", required=True)
            p.add_argument("--metrics")
            p.add_argument("--out", required=True)
        if name == "patch":
            p.add_argument("--changes", required=True)
            p.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    path = Path(a.file)
    if not path.exists():
        raise SystemExit(f"no such file: {path}")
    if a.cmd == "review":
        findings = json.loads(Path(a.findings).read_text(encoding="utf-8"))
        metrics = json.loads(Path(a.metrics).read_text(encoding="utf-8")) if a.metrics else None
        write_review(path, findings, metrics, Path(a.out))
        print(f"wrote {a.out}")
        return
    if a.cmd == "patch":
        for w in patch(path, json.loads(Path(a.changes).read_text(encoding="utf-8")), Path(a.out)):
            print("warning:", w)
        print(f"wrote {a.out} (original untouched)")
        return
    b = Book(path)
    if a.cmd == "inspect":
        print(json.dumps(inspect_book(b), indent=1, ensure_ascii=False, default=str))
    elif a.cmd == "dump":
        for s in ([a.sheet] if a.sheet else b.sheets):
            print(dump(b, s, a.range, a.max_rows), "\n")
    elif a.cmd == "extract":
        tables = [t for s in ([a.sheet] if a.sheet else b.sheets) for t in find_tables(b, s)]
        if a.json:
            print(json.dumps(tables, indent=1, ensure_ascii=False, default=str))
        else:
            for t in tables:
                print(f"## {t['sheet']} (header row {t['header_row']}): {', '.join(t['periods'])}")
                for lab, it in t["items"].items():
                    vals = ", ".join("" if v is None else f"{v:,.4g}" for v in it["values"])
                    print(f"  {lab[:40]:40} -> {it['canonical'] or '-':22} {vals}")
                print()
            if not tables:
                print("no period-by-line-item tables found; use dump to read the sheet")
    elif a.cmd == "ratios":
        print(json.dumps(ratios(b, a.sheet), indent=1, ensure_ascii=False, default=str))
    elif a.cmd == "audit":
        fs = audit(b)
        if a.json:
            print(json.dumps(fs, indent=1, ensure_ascii=False, default=str))
        else:
            for f in fs:
                print(f"[{f['severity']:6}] {f['sheet']}!{f['cell']}  {f['check']}: {f['detail']}"
                      + (f"\n          fix: {f['fix']}" if f['fix'] else ""))
            print(f"\n{len(fs)} findings")


if __name__ == "__main__":
    main(sys.argv[1:])
